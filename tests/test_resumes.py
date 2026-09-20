from __future__ import annotations

import datetime
import io
import tempfile
from pathlib import Path

import docx
import pymupdf
import pytest
from sqlalchemy.orm import Session

from pyjobs.models import JobApplication, Resume, UserPreference
from pyjobs.services.resume_parser import (
    extract_text_from_docx,
    extract_text_from_pdf,
    generate_download_filename,
    score_resume_match,
)


@pytest.fixture(autouse=True)
def temp_upload_dir(client):
    with tempfile.TemporaryDirectory() as tmpdir:
        client.app.state.upload_dir = Path(tmpdir)
        yield Path(tmpdir)
        client.app.state.upload_dir = None


def make_dummy_pdf(text: str = "Senior Python Developer with FastAPI and Pytest skills") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((54, 100), text, fontsize=12)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def make_dummy_docx(text: str = "Full Stack Engineer proficient in React and Python") -> bytes:
    doc = docx.Document()
    doc.add_heading("Resume", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_create_resume_with_pdf(client, db_session: Session, temp_upload_dir: Path):
    pdf_bytes = make_dummy_pdf("Expert in Python, Cloud, and Microservices")
    response = client.post(
        "/resumes",
        data={
            "title": "Backend Architect",
            "tags": "Python, Cloud, Backend",
            "description": "Targeting principal backend roles",
            "change_notes": "Initial PDF upload",
        },
        files={"file": ("my_resume.pdf", pdf_bytes, "application/pdf")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert response.headers["location"] == "/resumes"

    resume = db_session.query(Resume).filter(Resume.title == "Backend Architect").first()
    assert resume is not None
    assert resume.tags == "Python, Cloud, Backend"
    assert len(resume.versions) == 1

    v1 = resume.versions[0]
    assert v1.version_number == 1
    assert v1.file_type == "pdf"
    assert v1.original_filename == "my_resume.pdf"
    assert "Expert in Python" in (v1.extracted_text or "")
    assert Path(v1.file_path).exists()
    assert Path(v1.pdf_path).exists()
    assert str(temp_upload_dir) in v1.file_path


def test_create_resume_with_docx(client, db_session: Session, temp_upload_dir: Path):
    docx_bytes = make_dummy_docx("Full Stack Engineer with 7 years building SaaS")
    response = client.post(
        "/resumes",
        data={
            "title": "Full Stack Lead",
            "tags": "Full Stack, React, Python",
            "description": "Frontend & Backend roles",
            "change_notes": "First Word draft",
        },
        files={
            "file": (
                "lead_resume.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert response.status_code == 303

    resume = db_session.query(Resume).filter(Resume.title == "Full Stack Lead").first()
    assert resume is not None
    assert len(resume.versions) == 1

    v1 = resume.versions[0]
    assert v1.file_type == "docx"
    assert "Full Stack Engineer with 7 years" in (v1.extracted_text or "")
    # Check that a PDF was generated for the DOCX
    assert Path(v1.file_path).exists()
    assert Path(v1.pdf_path).exists()
    assert v1.pdf_path.endswith(".pdf")


def test_upload_subsequent_version(client, db_session: Session):
    pdf_v1 = make_dummy_pdf("Version 1 content")
    client.post(
        "/resumes",
        data={"title": "Data Scientist", "tags": "ML, Python"},
        files={"file": ("v1.pdf", pdf_v1, "application/pdf")},
    )
    resume = db_session.query(Resume).filter(Resume.title == "Data Scientist").first()
    assert resume is not None
    assert len(resume.versions) == 1

    pdf_v2 = make_dummy_pdf("Version 2 updated with PyTorch and LLMs")
    resp_v2 = client.post(
        f"/resumes/{resume.id}/versions",
        data={"change_notes": "Added PyTorch and LLM achievements"},
        files={"file": ("v2.pdf", pdf_v2, "application/pdf")},
        follow_redirects=False,
    )
    assert resp_v2.status_code == 303

    db_session.refresh(resume)
    assert len(resume.versions) == 2
    # Versions ordered descending by version_number
    assert resume.versions[0].version_number == 2
    assert resume.versions[0].change_notes == "Added PyTorch and LLM achievements"
    assert "PyTorch and LLMs" in (resume.versions[0].extracted_text or "")
    assert resume.versions[1].version_number == 1


def test_inline_pdf_view_endpoint(client, db_session: Session):
    pdf_bytes = make_dummy_pdf("Content for inline PDF viewing")
    client.post(
        "/resumes",
        data={"title": "DevOps Engineer"},
        files={"file": ("devops.pdf", pdf_bytes, "application/pdf")},
    )
    resume = db_session.query(Resume).filter(Resume.title == "DevOps Engineer").first()
    assert resume is not None
    v1 = resume.versions[0]

    resp = client.get(f"/resumes/versions/{v1.id}/view")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert "inline" in resp.headers["content-disposition"]


def test_download_filename_customization_and_dating(client, db_session: Session):
    # Set preferences
    pref = db_session.query(UserPreference).first()
    if not pref:
        pref = UserPreference()
        db_session.add(pref)
    pref.candidate_name = "Marissa G"
    pref.resume_filename_pattern = "{name} {date}.{ext}"
    pref.resume_date_format = "%m-%d-%Y"
    db_session.commit()

    docx_bytes = make_dummy_docx("Word resume content for retrieval testing")
    client.post(
        "/resumes",
        data={"title": "Product Manager"},
        files={
            "file": (
                "pm_resume.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    resume = db_session.query(Resume).filter(Resume.title == "Product Manager").first()
    assert resume is not None
    v1 = resume.versions[0]

    today_str = datetime.date.today().strftime("%m-%d-%Y")
    expected_pdf_name = f"Marissa G {today_str}.pdf"

    import urllib.parse

    # Test PDF download (rendered from docx)
    resp_pdf = client.get(f"/resumes/versions/{v1.id}/download?format=pdf")
    assert resp_pdf.status_code == 200
    assert resp_pdf.headers["content-type"] == "application/pdf"
    disposition_pdf = urllib.parse.unquote(resp_pdf.headers["content-disposition"])
    assert expected_pdf_name in disposition_pdf

    # Test Original DOCX download
    expected_docx_name = f"Marissa G {today_str}.docx"
    resp_docx = client.get(f"/resumes/versions/{v1.id}/download?format=original")
    assert resp_docx.status_code == 200
    disposition_docx = urllib.parse.unquote(resp_docx.headers["content-disposition"])
    assert expected_docx_name in disposition_docx

    # Test custom filename override
    resp_custom = client.get(
        f"/resumes/versions/{v1.id}/download?format=pdf&custom_filename=Marissa_G_Google_Application"
    )
    assert resp_custom.status_code == 200
    disposition_custom = urllib.parse.unquote(resp_custom.headers["content-disposition"])
    assert "Marissa_G_Google_Application.pdf" in disposition_custom


def test_tag_filtering_and_search(client, db_session: Session):
    pdf1 = make_dummy_pdf("Python FastAPI")
    pdf2 = make_dummy_pdf("Frontend React TypeScript")
    client.post(
        "/resumes",
        data={"title": "Python Developer", "tags": "Python, Backend"},
        files={"file": ("py.pdf", pdf1, "application/pdf")},
    )
    client.post(
        "/resumes",
        data={"title": "React Specialist", "tags": "Frontend, React"},
        files={"file": ("react.pdf", pdf2, "application/pdf")},
    )

    resp_all = client.get("/resumes")
    assert "Python Developer" in resp_all.text
    assert "React Specialist" in resp_all.text

    resp_tag = client.get("/resumes?tag=Python")
    assert "Python Developer" in resp_tag.text
    assert "React Specialist" not in resp_tag.text

    resp_search = client.get("/resumes?q=Specialist")
    assert "React Specialist" in resp_search.text
    assert "Python Developer" not in resp_search.text


def test_application_resume_link(client, db_session: Session):
    pdf = make_dummy_pdf("Application test")
    client.post(
        "/resumes",
        data={"title": "Lead Software Engineer"},
        files={"file": ("lead.pdf", pdf, "application/pdf")},
    )
    resume = db_session.query(Resume).filter(Resume.title == "Lead Software Engineer").first()
    assert resume is not None
    v1 = resume.versions[0]

    # Create application linked to resume version
    resp = client.post(
        "/applications",
        data={
            "company": "Stripe",
            "title": "Lead Software Engineer",
            "status": "applied",
            "resume_version_id": v1.id,
        },
        follow_redirects=False,
    )
    assert resp.status_code == 303

    app_record = db_session.query(JobApplication).filter(JobApplication.company == "Stripe").first()
    assert app_record is not None
    assert app_record.resume_version_id == v1.id
    assert app_record.resume_version is not None
    assert app_record.resume_version.resume.title == "Lead Software Engineer"

    # Check detail page displays linked resume
    resp_detail = client.get(f"/applications/{app_record.id}")
    assert resp_detail.status_code == 200
    assert "Lead Software Engineer" in resp_detail.text
    assert f"/resumes/versions/{v1.id}/view" in resp_detail.text


def test_delete_version_and_resume_removes_files(client, db_session: Session):
    pdf1 = make_dummy_pdf("V1")
    client.post(
        "/resumes",
        data={"title": "Ephemeral Resume"},
        files={"file": ("v1.pdf", pdf1, "application/pdf")},
    )
    resume = db_session.query(Resume).filter(Resume.title == "Ephemeral Resume").first()
    assert resume is not None
    v1_path = Path(resume.versions[0].file_path)
    assert v1_path.exists()
    resume_id = resume.id

    pdf2 = make_dummy_pdf("V2")
    client.post(
        f"/resumes/{resume_id}/versions",
        data={"change_notes": "V2"},
        files={"file": ("v2.pdf", pdf2, "application/pdf")},
    )
    db_session.refresh(resume)
    assert resume is not None
    v2_id = resume.versions[0].id
    v2_path = Path(resume.versions[0].file_path)
    assert v2_path.exists()

    # Delete single version v2
    del_v2 = client.delete(f"/resumes/versions/{v2_id}")
    assert del_v2.status_code == 200
    assert not v2_path.exists()
    assert v1_path.exists()

    # Delete entire resume
    del_resume = client.delete(f"/resumes/{resume_id}")
    assert del_resume.status_code == 200
    assert not v1_path.exists()
    assert db_session.query(Resume).filter(Resume.id == resume_id).first() is None


def test_invalid_file_type_rejected(client):
    resp = client.post(
        "/resumes",
        data={"title": "Invalid Resume"},
        files={"file": ("script.sh", b"#!/bin/bash\necho bad", "text/plain")},
    )
    assert resp.status_code == 400
    assert "Unsupported file format" in resp.text


def test_resume_parser_pure_functions():
    pdf_bytes = make_dummy_pdf("Testing pure PDF parser functionality")
    text_pdf = extract_text_from_pdf(pdf_bytes)
    assert "Testing pure PDF parser" in text_pdf

    docx_bytes = make_dummy_docx("Testing pure DOCX parser functionality")
    text_docx = extract_text_from_docx(docx_bytes)
    assert "Testing pure DOCX parser" in text_docx

    # Filename generation testing
    fname = generate_download_filename(
        pattern="{name} - {title} - {date}.{ext}",
        date_format="%Y-%m-%d",
        candidate_name="Jane Doe",
        resume_title="Backend Developer",
        version_number=1,
        file_ext="pdf",
        retrieval_date=datetime.date(2026, 9, 10),
    )
    assert fname == "Jane Doe - Backend Developer - 2026-09-10.pdf"

    # Match scoring testing
    score_high = score_resume_match(
        resume_title="Senior Python Backend Developer",
        resume_tags="Python, FastAPI, Postgres, Docker",
        job_title="Senior Python Backend Engineer",
        job_description=(
            "Seeking a backend engineer with strong Python, FastAPI, and Postgres experience."
        ),
    )
    assert score_high >= 0.5

    score_low = score_resume_match(
        resume_title="Senior Python Backend Developer",
        resume_tags="Python, FastAPI",
        job_title="Registered Nurse - Emergency Room",
        job_description="Patient care in hospital environment.",
    )
    assert score_low < 0.1


def test_create_and_edit_resume_person(client, db_session: Session):
    pdf = make_dummy_pdf("Candidate profile")
    resp = client.post(
        "/resumes",
        data={
            "title": "Douglas - MSAT Focused",
            "person": "Douglas Jaymes Caristo",
            "tags": "MSAT Engineer, CAPA",
            "description": "Biotech MSAT specialist",
        },
        files={"file": ("douglas.pdf", pdf, "application/pdf")},
        follow_redirects=False,
    )
    assert resp.status_code == 303

    resume = db_session.query(Resume).filter(Resume.title == "Douglas - MSAT Focused").first()
    assert resume is not None
    assert resume.person == "Douglas Jaymes Caristo"

    # Verify displayed on /resumes
    list_resp = client.get("/resumes")
    assert list_resp.status_code == 200
    assert "Douglas Jaymes Caristo" in list_resp.text
    assert "Filter by Person:" in list_resp.text

    # Edit person
    edit_resp = client.post(
        f"/resumes/{resume.id}/edit",
        data={
            "title": "Douglas - Senior MSAT",
            "person": "Douglas J. Caristo",
            "tags": "MSAT, Validation",
            "description": "Updated focus",
        },
        follow_redirects=False,
    )
    assert edit_resp.status_code == 303
    db_session.refresh(resume)
    assert resume.title == "Douglas - Senior MSAT"
    assert resume.person == "Douglas J. Caristo"
    assert resume.tags == "MSAT, Validation"


def test_filter_resumes_by_person(client, db_session: Session):
    pdf = make_dummy_pdf("Content")
    client.post(
        "/resumes",
        data={"title": "Douglas Profile", "person": "Douglas Caristo"},
        files={"file": ("doug.pdf", pdf, "application/pdf")},
    )
    client.post(
        "/resumes",
        data={"title": "Marissa Profile", "person": "Marissa Gaeta"},
        files={"file": ("marissa.pdf", pdf, "application/pdf")},
    )

    # Filter for Douglas
    resp_doug = client.get("/resumes?person=Douglas%20Caristo")
    assert resp_doug.status_code == 200
    assert "Douglas Profile" in resp_doug.text
    assert "Marissa Profile" not in resp_doug.text

    # Filter for Marissa
    resp_marissa = client.get("/resumes?person=Marissa%20Gaeta")
    assert resp_marissa.status_code == 200
    assert "Marissa Profile" in resp_marissa.text
    assert "Douglas Profile" not in resp_marissa.text

    # Filter all
    resp_all = client.get("/resumes?person=all")
    assert resp_all.status_code == 200
    assert "Douglas Profile" in resp_all.text
    assert "Marissa Profile" in resp_all.text


def test_download_filename_maps_name_to_person_with_fallback(client, db_session: Session):
    import urllib.parse

    pref = db_session.query(UserPreference).first()
    if not pref:
        pref = UserPreference()
        db_session.add(pref)
    pref.candidate_name = "Global Fallback"
    pref.resume_filename_pattern = "{name} {date}.{ext}"
    pref.resume_date_format = "%m-%d-%Y"
    db_session.commit()

    pdf = make_dummy_pdf("Resume file")
    # 1. Resume WITH person specified
    client.post(
        "/resumes",
        data={"title": "MSAT Resume", "person": "Douglas Jaymes Caristo"},
        files={"file": ("doc.pdf", pdf, "application/pdf")},
    )
    res_with_person = (
        db_session.query(Resume).filter(Resume.person == "Douglas Jaymes Caristo").first()
    )
    assert res_with_person is not None
    v_person = res_with_person.versions[0]

    resp_download_person = client.get(f"/resumes/versions/{v_person.id}/download")
    assert resp_download_person.status_code == 200
    disp_person = urllib.parse.unquote(resp_download_person.headers["content-disposition"])
    assert "Douglas Jaymes Caristo" in disp_person
    assert "Global Fallback" not in disp_person

    # 2. Resume WITHOUT person specified -> falls back to global candidate_name
    client.post(
        "/resumes",
        data={"title": "Anonymous Resume", "person": ""},
        files={"file": ("anon.pdf", pdf, "application/pdf")},
    )
    res_anon = db_session.query(Resume).filter(Resume.title == "Anonymous Resume").first()
    assert res_anon is not None
    v_anon = res_anon.versions[0]

    resp_download_anon = client.get(f"/resumes/versions/{v_anon.id}/download")
    assert resp_download_anon.status_code == 200
    disp_anon = urllib.parse.unquote(resp_download_anon.headers["content-disposition"])
    assert "Global Fallback" in disp_anon
